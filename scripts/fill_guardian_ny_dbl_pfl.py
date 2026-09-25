#!/usr/bin/env python3
"""Fill Guardian's NY DBL + PFL quarterly premium report from a payroll export.

Inputs:
  - the Guardian "PREMIUM REPORT - QUARTERLY" PDF for the billing period
    (a blank statement, or a previously filled one -- any flattened fill
    layer is stripped before writing new values)
  - the QuickBooks payroll report .xlsx: one row per employee, already
    summed across every pay run in the billing period. Columns used:
      Employee, Employee gross pay, Start date (= HIRE DATE),
      Identified gender (falls back to Legal gender),
      Year-to-date gross pay (optional; any header containing "year to
      date"/"YTD" and "gross"). YTD must be as of the period end, i.e.
      include this period's pay. It is used to apply the PFL annual cap:
      covered = min(period wages, cap - (YTD - period wages)). Without it,
      the cap is applied to period wages only and a warning is printed.

Everything that changes quarter to quarter (period dates, DBL rates,
minimum premium, DBL and PFL wage caps, PFL rate) is read from the PDF
text rather than hardcoded.

Usage:
  python3 scripts/fill_guardian_ny_dbl_pfl.py STATEMENT.pdf PAYROLL.xlsx \
      --out filled.pdf --preparer "Name" --email "addr@example.com" \
      [--preview preview.png]
"""
import argparse
import calendar
import datetime as dt
import io
import re
import sys
from decimal import Decimal, ROUND_HALF_UP

import openpyxl
import pypdf
from pypdf.generic import ContentStream, NameObject
from reportlab.pdfgen import canvas

D = Decimal

# Bottom-left corners (PDF points) of the entry boxes on Guardian form
# NYGQSEP. Text is drawn 3.31pt above, Helvetica 10, matching the
# form's own field appearance.
MONTH_COL_X = [94.0, 174.96, 255.6]  # 1st, 2nd, 3rd month of the quarter
POS = {
    'stat_payroll': (498.96, 469.92),
    'm_total': (325.92, 395.28), 'm_prem': (504.24, 393.6),
    'f_total': (325.92, 370.56), 'f_prem': (504.24, 372.24),
    'A': (504, 348.48), 'B': (504, 318.96),
    'pfl_m': (78, 191.28), 'pfl_f': (78, 163.2),
    'C': (504.24, 216.72), 'E': (504.24, 165.12), 'F': (504.24, 144.96),
    'preparer': (84.72, 122.16), 'email': (363.84, 122.16),
}
M_ROW_Y, F_ROW_Y = 393.6, 370.56


def cents(x):
    return D(x).quantize(D('0.01'), ROUND_HALF_UP)


def money(x):
    return f"${cents(x):,.2f}"


def num(s):
    return D(s.replace(',', ''))


def parse_statement(reader):
    text = reader.pages[0].extract_text(extraction_mode='layout')
    flat = re.sub(r'\s+', ' ', text)

    def find(pattern, label):
        m = re.search(pattern, flat)
        if not m:
            sys.exit(f"Could not find {label} on the statement; check the form layout.")
        return m

    dates = find(r'(\d\d/\d\d/\d{4}) (\d\d/\d\d/\d{4}) (\d\d/\d\d/\d{4})', 'period dates')
    to_date = lambda s: dt.datetime.strptime(s, '%m/%d/%Y').date()
    months = find(r'MONTH\s+([A-Z]{3})\s+([A-Z]{3})\s+([A-Z]{3})', 'month headers')
    return {
        'policy': find(r'(\d{8}-\d{4})', 'policy number').group(1),
        'begin': to_date(dates.group(1)),
        'end': to_date(dates.group(2)),
        'due': to_date(dates.group(3)),
        'month_names': list(months.groups()),
        'rate_m': num(find(r'MALES.*?x \$\s*([\d.]+)', 'male DBL rate').group(1)),
        'rate_f': num(find(r'FEMALES.*?x \$\s*([\d.]+)', 'female DBL rate').group(1)),
        'min_prem': num(find(r'Minimum Quarterly Premium: \$\s*([\d.,]+)', 'minimum premium').group(1)),
        'dbl_cap': num(find(r'MAXIMUM OF \$([\d,]+\.\d\d)', 'DBL quarterly wage cap').group(1)),
        'pfl_cap': num(find(r'not to exceed \$([\d,]+\.\d\d)', 'PFL annual wage cap').group(1)),
        'pfl_rate': num(find(r'Premium Rate - [\d.]+% of Wages: ([\d.]+)', 'PFL rate').group(1)),
    }


def quarter_months(st):
    """The three months named on the form, as (name, first_day, last_day)."""
    year = st['end'].year
    out = []
    for name in st['month_names']:
        m = dt.datetime.strptime(name.title(), '%b').month
        y = year - 1 if m > st['end'].month else year
        out.append((name, dt.date(y, m, 1), dt.date(y, m, calendar.monthrange(y, m)[1])))
    return out


def load_payroll(path):
    ws = openpyxl.load_workbook(path, data_only=True).active
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    col = {}
    for i, h in enumerate(header):
        col.setdefault(h, i)  # first occurrence wins ("Employee" appears twice)
    need = ['Employee', 'Employee gross pay', 'Start date']
    missing = [c for c in need if c not in col]
    if missing:
        sys.exit(f"Payroll report is missing columns: {missing}")
    gcol = col.get('Identified gender', col.get('Legal gender'))
    ycol = next((i for h, i in col.items() if h and 'gross' in str(h).lower()
                 and re.search(r'year.to.date|\bytd\b', str(h).lower())), None)
    emps = []
    for r in rows:
        name = r[col['Employee']]
        if not name or name == 'All' or r[col['Employee gross pay']] is None:
            continue
        hired = r[col['Start date']]
        hired = hired.date() if isinstance(hired, dt.datetime) else hired
        gender = (r[gcol] or '').strip().title() if gcol is not None else ''
        if gender not in ('Male', 'Female'):
            sys.exit(f"{name}: gender is {gender!r}; the form only has Male/Female rows.")
        ytd = r[ycol] if ycol is not None else None
        emps.append({'name': name, 'gender': gender, 'hired': hired,
                     'wages': D(str(r[col['Employee gross pay']])),
                     'ytd': D(str(ytd)) if ytd is not None else None})
    return emps, ycol is not None


def strip_prior_fill(writer):
    """Remove flattened field layers (form XObjects wrapping /Tx content)."""
    page = writer.pages[0]
    xobjs = page['/Resources'].get('/XObject') or {}

    def has_fields(x):
        x = x.get_object()
        if b'/Tx BMC' in x.get_data():
            return True
        sub = (x.get('/Resources') or {}).get('/XObject') or {}
        return any(has_fields(v) for v in sub.values())

    drop = {NameObject(k) for k, v in xobjs.items()
            if v.get_object().get('/Subtype') == '/Form' and has_fields(v)}
    if drop:
        cs = ContentStream(page.get_contents(), writer)
        cs.operations = [op for op in cs.operations
                         if not (op[1] == b'Do' and op[0][0] in drop)]
        page.replace_contents(cs)
    return len(drop)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('statement')
    ap.add_argument('payroll')
    ap.add_argument('--out', required=True)
    ap.add_argument('--preparer', default='')
    ap.add_argument('--email', default='')
    ap.add_argument('--preview', help='optional PNG render of page 1 (needs pypdfium2)')
    a = ap.parse_args()

    reader = pypdf.PdfReader(a.statement)
    st = parse_statement(reader)
    emps, has_ytd = load_payroll(a.payroll)
    months = quarter_months(st)

    # A month is on the form only if it overlaps the billing period
    # (earlier months are pre-printed N/A). An employee counts in a month
    # if hired on or before that month's last day.
    counts = []  # (col_index, name, males, females) for active months
    for i, (name, first, last) in enumerate(months):
        if last < st['begin'] or first > st['end']:
            continue
        hired = [e for e in emps if e['hired'] <= last]
        counts.append((i, name, sum(e['gender'] == 'Male' for e in hired),
                       sum(e['gender'] == 'Female' for e in hired)))

    m_tot = sum(c[2] for c in counts)
    f_tot = sum(c[3] for c in counts)
    m_prem = cents(m_tot * st['rate_m'])
    f_prem = cents(f_tot * st['rate_f'])
    A = m_prem + f_prem
    B = max(A, st['min_prem'])

    in_period = [e for e in emps if e['hired'] <= st['end']]
    stat = sum(min(e['wages'], st['dbl_cap']) for e in in_period)
    for e in emps:
        prior = (e['ytd'] - e['wages']) if e['ytd'] is not None else D(0)
        if prior < 0:
            sys.exit(f"{e['name']}: YTD gross {e['ytd']} is less than period gross {e['wages']}; "
                     "YTD must be as of the period end.")
        e['pfl_wages'] = max(D(0), min(e['wages'], st['pfl_cap'] - prior))
    C = sum(e['pfl_wages'] for e in in_period)
    E = cents(C * st['pfl_rate'])
    F = B + E
    pfl_m = sum(e['gender'] == 'Male' for e in in_period)
    pfl_f = sum(e['gender'] == 'Female' for e in in_period)

    vals = {
        'stat_payroll': money(stat),
        'm_total': m_tot, 'm_prem': money(m_prem),
        'f_total': f_tot, 'f_prem': money(f_prem),
        'A': money(A), 'B': money(B),
        'pfl_m': pfl_m, 'pfl_f': pfl_f,
        'C': money(C), 'E': money(E), 'F': money(F),
        'preparer': a.preparer, 'email': a.email,
    }

    buf = io.BytesIO()
    cv = canvas.Canvas(buf, pagesize=(612, 792))
    cv.setFont('Helvetica', 10)
    for key, v in vals.items():
        x, y = POS[key]
        cv.drawString(x, y + 3.31, str(v))
    for i, _, m, f in counts:
        cv.drawString(MONTH_COL_X[i], M_ROW_Y + 3.31, str(m))
        cv.drawString(MONTH_COL_X[i], F_ROW_Y + 3.31, str(f))
    cv.save()
    buf.seek(0)

    writer = pypdf.PdfWriter()
    writer.append(reader)
    stripped = strip_prior_fill(writer)
    writer.pages[0].merge_page(pypdf.PdfReader(buf).pages[0])
    writer.write(a.out)

    if a.preview:
        import pypdfium2
        pypdfium2.PdfDocument(a.out)[0].render(scale=2).to_pil().save(a.preview)

    print(f"Policy {st['policy']}  period {st['begin']:%m/%d/%Y}-{st['end']:%m/%d/%Y}  due {st['due']:%m/%d/%Y}")
    print(f"Rates: M ${st['rate_m']}  F ${st['rate_f']}  min ${st['min_prem']}  "
          f"DBL cap ${st['dbl_cap']:,}/qtr  PFL cap ${st['pfl_cap']:,}/yr  PFL rate {st['pfl_rate']}")
    if stripped:
        print(f"Stripped {stripped} prior fill layer(s) from the statement.")
    print("\nEmployees:")
    for e in emps:
        flag = '' if e['hired'] <= st['end'] else '  (hired after period -- excluded)'
        ytd = f"  YTD {money(e['ytd']):>12}" if e['ytd'] is not None else ''
        print(f"  {e['name']:<25} {e['gender']:<6} hired {e['hired']:%m/%d/%Y}  wages {money(e['wages']):>12}{ytd}{flag}")
        if e['pfl_wages'] < e['wages']:
            print(f"    ! PFL annual cap reached; covered wages {money(e['pfl_wages'])}")
    if not has_ytd:
        print("  ! No year-to-date gross column found: PFL annual cap applied to period wages only.")
    print("\nNY DBL headcount by month:")
    for _, name, m, f in counts:
        print(f"  {name}: {m} male, {f} female")
    print(f"  Males   {m_tot} x ${st['rate_m']} = {money(m_prem)}")
    print(f"  Females {f_tot} x ${st['rate_f']} = {money(f_prem)}")
    print(f"  A = {money(A)}   B = {money(B)}   DBL statutory payroll = {money(stat)}")
    print(f"NY PFL: {pfl_m} male, {pfl_f} female   C = {money(C)}   E = {money(E)}")
    print(f"F. TOTAL DUE = {money(F)}")
    print(f"\nWrote {a.out}")


if __name__ == '__main__':
    main()
